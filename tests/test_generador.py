from datetime import date, timedelta

from docx import Document

from generador import Generador
from interfaz import App
from utils import validar_fechas_no_futuras


def test_permite_fechas_iguales_o_anteriores_a_hoy():
    hoy = date.today()

    assert validar_fechas_no_futuras(hoy, [hoy, hoy - timedelta(days=1)])


def test_rechaza_cualquier_fecha_futura():
    hoy = date.today()

    assert not validar_fechas_no_futuras(hoy + timedelta(days=1), [hoy])
    assert not validar_fechas_no_futuras(hoy, [hoy + timedelta(days=1)])


def test_reemplazar_desde_marca_parrafo_para_relleno_con_igual(tmp_path):
    ruta_docx = tmp_path / "plantilla.docx"
    doc = Document()
    doc.add_paragraph("texto base")
    doc.save(ruta_docx)

    generador = Generador(str(ruta_docx))
    generador.reemplazar_desde(
        "texto base",
        [[("texto nuevo", True)]],
        con_lider_final=True,
    )

    assert generador._indices_pendientes_iguales
    assert generador._indices_pendientes_iguales[-1][1] == "="


def test_linea_de_relleno_no_expone_marcador_interno(tmp_path):
    ruta_docx = tmp_path / "plantilla.docx"
    Document().save(ruta_docx)

    generador = Generador(str(ruta_docx))
    generador._agregar_parrafo_con_lider(
        {"runs": [("*", False)], "fill_char": "*"}
    )

    assert generador.doc.paragraphs[-1].text == "*"
    assert "runsfill_char" not in generador.doc.paragraphs[-1].text


def test_extrae_bloque_fe_y_testado_desde_documento_original(tmp_path):
    ruta_docx = tmp_path / "plantilla.docx"
    doc = Document()
    doc.add_paragraph(
        "FE DE CONTENIDO Y LECTURA: INSTRUIDOS LOS OTORGANTES DEL CONTENIDO DEL PRESENTE "
        "INSTRUMENTO POR LA LECTURA QUE LES HIZO EL NOTARIO, SE RATIFICAN EN SU CONTENIDO, "
        "PROCEDIENDO A FIRMAR JUNTO CONMIGO, DE LO QUE DOY FE. LA PRESENTE ESCRITURA SE INICIA EN LA FOJA SERIE B 5903304 V Y TERMINA EN LA FOJA SERIE B 5903308 V. DOY FE."
    )
    doc.add_paragraph("TESTADO: VAMO A CORREGIR TA VAINA")
    doc.add_paragraph("EL PROCESO DE FIRMAS CONCLUYO")
    doc.save(ruta_docx)

    generador = Generador(str(ruta_docx))
    bloque = App._extraer_bloque_fe_y_testado(generador)

    assert "FE DE CONTENIDO Y LECTURA" in bloque
    assert "DOY FE." in bloque
    assert bloque.count("DOY FE") >= 2
    assert "TESTADO: VAMO A CORREGIR TA VAINA" in bloque


def test_extrae_bloque_fe_original_hasta_segundo_doy_fe_y_testado(tmp_path):
    ruta_docx = tmp_path / "plantilla.docx"
    doc = Document()
    doc.add_paragraph(
        "FE DE CONTENIDO Y LECTURA: INSTRUIDOS LOS OTORGANTES DEL CONTENIDO DEL PRESENTE "
        "INSTRUMENTO POR LA LECTURA QUE LES HIZO EL NOTARIO, SE RATIFICAN EN SU CONTENIDO, "
        "PROCEDIENDO A FIRMAR JUNTO CONMIGO, DE LO QUE DOY FE. LA PRESENTE ESCRITURA SE INICIA EN LA FOJA SERIE B 5903304 V Y TERMINA EN LA FOJA SERIE B 5903308 V. DOY FE."
    )
    doc.add_paragraph("TESTADO: VAMO A CORREGIR TA VAINA")
    doc.add_paragraph("EL PROCESO DE FIRMAS CONCLUYO")
    doc.save(ruta_docx)

    generador = Generador(str(ruta_docx))
    bloque = App._extraer_bloque_fe_y_testado(generador)

    assert "FE DE CONTENIDO Y LECTURA" in bloque
    assert "DOY FE." in bloque
    assert bloque.count("DOY FE") >= 2
    assert "TESTADO: VAMO A CORREGIR TA VAINA" in bloque
